"""Extract performance data from the RBC workbooks.

Usage: python build/parse_perf.py <folder with the workbooks>
Writes data/performance.json:
  hpr      period returns since inception, from the latest "HPR since inception" workbook
           (first sheet): {start, rows:[{end, months, smpt, tsx, sp}]}, returns as decimals.
           tsx and sp are total-return index changes; smpt is RBC's gross HPR.
  periods  one entry per "Portfolio Holdings" workbook, from sheets 06 (beta), 08 (HPR)
           and 09 (return compared to index).
"""
import glob
import json
import os
import re
import sys
from datetime import datetime

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if len(sys.argv) < 2:
    sys.exit('Usage: python build/parse_perf.py <folder with the workbooks>')
SRC = sys.argv[1]


def dated(pattern):
    """Workbooks matching pattern whose name starts with YYYY MM DD, oldest first."""
    out = []
    for f in sorted(glob.glob(os.path.join(SRC, pattern))):
        name = os.path.basename(f)
        m = re.match(r'(\d{4}) (\d{2}) (\d{2})', name)
        if m and not name.startswith('~$'):
            out.append(('-'.join(m.groups()), f))
    return out


def num(c):
    return c if isinstance(c, (int, float)) else None


def cells(r):
    return [c for c in r if c is not None]


def label(r):
    c = cells(r)
    return str(c[0]).strip().lower() if c else ''


def parse_hpr(path):
    """Rows are: year, period end, months, months since inception, SMPT rate (%), SMPT value,
    TSX level, TSX rate, TSX value, S&P level, S&P rate, S&P value. The first row is the base."""
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
    start, out = None, []
    for r in ws.iter_rows(values_only=True):
        c = cells(r)
        if len(c) < 8 or not isinstance(c[1], datetime):
            continue
        if start is None:  # base row: no period return yet
            start = c[1].strftime('%Y-%m-%d')
            continue
        out.append(dict(end=c[1].strftime('%Y-%m-%d'), months=c[2], smpt=c[4] / 100, tsx=c[7], sp=c[10]))
    return dict(start=start, rows=out)


def parse_period(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = lambda key: next((w for w in wb.worksheets if key in w.title.lower()), None)
    p = {}

    # 09: per-class representative index returns, composite benchmark, HPR and abnormal return
    ws = sheet('compared to index')
    rows = [r for r in ws.iter_rows(values_only=True)]
    hi = next(i for i, r in enumerate(rows) if 'Price on' in [str(c).strip() for c in r if c is not None])
    sub = [str(c).strip() if c is not None else '' for c in rows[hi + 1]]
    dates = [c for c in rows[hi + 1] if isinstance(c, datetime)]
    p['start'], p['end'] = (d.strftime('%Y-%m-%d') for d in dates[:2])
    ci, wi, ii = sub.index('Asset Class'), sub.index('of Total'), sub.index('Representative Index')
    ri = next(i for i, s in enumerate(sub) if re.fullmatch(r'\d+ days', s))  # "HPR over N days"
    p['classes'] = []
    for r in rows[hi + 2:]:
        if label(r).startswith('total holdings'):
            p['bench'] = num(cells(r)[-1])
            break
        if isinstance(r[ci], str) and isinstance(r[ii], str) and num(r[wi]) is not None:
            p['classes'].append(dict(cls=r[ci].strip(), w=r[wi], index=r[ii].strip(), ret=num(r[ri]) or 0))
    for r in rows:
        lab, c = label(r), cells(r)
        if lab.startswith('holding period return'):
            p['hpr'] = num(c[1])
        elif lab.startswith('abnormal return'):
            p['abn'] = num(c[1])
        elif lab.startswith('annulaized abnormal') or lab.startswith('annualized abnormal'):
            p['abn_ann'] = num(c[1])

    # 08: gross HPR over the period
    ws = sheet('hpr')
    for r in ws.iter_rows(values_only=True):
        if re.match(r'\d+-day holding period', label(r)):
            p['hpr_gross'] = num(cells(r)[1])

    # 06: beta of the equity holdings (historic, expected). Skip sheets marked "Not Ready"
    # (copied from the previous period) and zero placeholders.
    ws = sheet('beta')
    p['beta'] = None
    if ws is not None and 'not ready' not in ws.title.lower():
        for r in ws.iter_rows(values_only=True):
            if label(r).startswith('beta of the equity'):
                b = [x for x in cells(r)[1:] if num(x)]
                if b:
                    p['beta'] = dict(hist=b[0], exp=b[1] if len(b) > 1 else None)
    return p


hpr_files = dated('*HPR since inception*.xlsx')
out = dict(hpr=parse_hpr(hpr_files[-1][1]) if hpr_files else dict(start=None, rows=[]),
           hpr_source=hpr_files[-1][0] if hpr_files else None,
           periods=[])
for date, f in dated('*Portfolio Holdings*.xlsx'):
    p = parse_period(f)
    p['snapshot'] = date
    out['periods'].append(p)
    print(date, '->', {k: (round(v, 4) if isinstance(v, float) else v) for k, v in p.items() if k != 'classes'},
          len(p['classes']), 'classes')
print('hpr:', len(out['hpr']['rows']), 'periods from', hpr_files[-1][0] if hpr_files else None)

json.dump(out, open(os.path.join(ROOT, 'data', 'performance.json'), 'w'), default=str, indent=1)
